from openpyxl import load_workbook
import os

def extrair_imagens(caminho_excel):

    planilha = load_workbook(caminho_excel)
    aba = planilha.active

    imagens = []

# Aqui _images pega a imagem que está em aba, no caso a principal, e a guarda, então dados guarda o valor dos dados da imagem obtidos em _data() que logo são inseridos em imagens os dados de cada imagem que lhe for passado como uma lista
    for imagem in aba._images:
        dados = imagem._data()
        imagens.append(dados)

    return imagens

if __name__ == "__main__":

    caminho = "temporario/Etiqueta_26-02187_Mateus.xlsx"

    imagens = extrair_imagens(caminho)

    print(f"Quantidade de imagens: {len(imagens)}")

    planilha = load_workbook(caminho)

    aba = planilha.active

# cria a pasta onde guardamos as imagens
    os.makedirs("temporario/imagens",exist_ok=True)

# O enumerate percorre as imagens e numer cada uma
    for i,imagem in enumerate(aba._images, start=1):

        dados = imagem._data()

        caminho = f"temporario/imagens/imagem_{i}.png"

# Cria um arquivo em modo de escrita binária, isso se dá porque transformamos nossa imagem em _data pra extrair seus dados
        with open(caminho, "wb") as arquivo:

# Grava a imagem no computador
            arquivo.write(dados)

        print(f"Imagem {i} salva em {caminho}")

    #     print(f"Imagem: {i}")
    #     print(f"Tamanho: {len(dados)} bytes")
    #     print(f"Posição: {imagem.anchor._from.col}, {imagem.anchor._from.row}")
    #     print(f"Dimensões: {imagem.width} x {imagem.height}")

    # print(f"\nCONFIGURAÇÃO DA PÁGINA")

    # print(f"Orientação: {aba.page_setup.orientation}")
    # print(f"Papel: {aba.page_setup.paperSize}")
    # print(f"Área de impressão: {aba.print_area}")

    # print(f'\nDIMENSÕES DAS COLUNAS')
    # for coluna, dimensao in aba.column_dimensions.items():
    #     print(coluna, dimensao.width)

    # print(f"\nALTURA DAS LINHAS")

    # for linha,dimensao in aba.row_dimensions.items():
    #     print(linha,dimensao.height)